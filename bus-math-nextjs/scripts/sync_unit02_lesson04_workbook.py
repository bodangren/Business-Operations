"""Copy the existing Lesson 4 source and complete its website reference copy."""

from __future__ import annotations

import argparse
import hashlib
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties


RESOURCES = Path(__file__).resolve().parents[1] / "public" / "resources"
SHEETS = [
    "Unadjusted TB", "Adjustments", "Adjusted TB", "Financial Statements",
    "Closing Entries", "Post-closing TB",
]


def prepare_website_copy(workbook) -> None:
    """Keep the source layout and make its labels and checks usable."""
    layouts = [(4, [1, 2, 23, 32, 33]), (5, [1, 2, 22]),
               (7, [1, 2, 22]), (3, [1, 2, 37]),
               (5, [1, 2, 17, 19]), (4, [1, 2, 18])]
    for sheet, (columns, note_rows) in zip(workbook.worksheets, layouts):
        for row in sheet.iter_rows():
            populated = [cell for cell in row if cell.value is not None]
            if row[0].row in note_rows and len(populated) == 1 and populated[0].column == 1:
                sheet.merge_cells(start_row=row[0].row, start_column=1,
                                  end_row=row[0].row, end_column=columns)
            for cell in populated:
                alignment = copy(cell.alignment)
                alignment.wrap_text = True
                cell.alignment = alignment
        sheet.row_dimensions[2].height = 36
        sheet.row_dimensions[4].height = 32
        sheet.column_dimensions["A"].width = 38
        if sheet.title == "Adjustments":
            sheet.column_dimensions["A"].width = 10
            for row in range(5, 17):
                alignment = copy(sheet.cell(row, 1).alignment)
                alignment.horizontal = "center"
                sheet.cell(row, 1).alignment = alignment
        for row in (25, 30):
            if sheet.title == "Unadjusted TB":
                sheet.row_dimensions[row].height = 32
        if sheet.title == "Financial Statements":
            sheet.row_dimensions[24].height = 32
            sheet.row_dimensions[37].height = 36
        sheet.page_setup.orientation = (
            "landscape" if sheet.title in ("Adjustments", "Adjusted TB", "Closing Entries") else "portrait"
        )
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_area = f"A1:{get_column_letter(columns)}{sheet.max_row}"

    workbook["Adjustments"]["D20"] = (
        '=IF(OR(COUNTA(C5:C16)<>12,COUNT(D5:E16)<>24),'
        '"Not finished",IF(D19=0,"Balanced","Review"))'
    )
    workbook["Adjusted TB"]["F21"] = '=IF(COUNT(D5:G19)<>60,"",F20-G20)'
    workbook["Closing Entries"]["D16"] = (
        '=IF(OR(COUNTA(C5:C13)<>9,COUNT(D5:E13)<>18),"",D15-E15)'
    )
    workbook["Post-closing TB"]["B16"] = '=IF(COUNT(B5:C14)<>20,"",B15-C15)'
    workbook["Unadjusted TB"]["A33"] = 'https://bodangren.github.io/Business-Operations/student/unit02/lesson05/#do'
    workbook.calculation = CalcProperties(fullCalcOnLoad=True, forceFullCalc=True, calcMode="auto")


def complete_reference(workbook) -> None:
    """Complete the existing manual worksheet with linked accounting results."""
    journal = workbook["Adjustments"]
    adjustments = (
        ("Supplies Expense", "Supplies", "='Unadjusted TB'!B7-'Unadjusted TB'!B24"),
        ("Insurance Expense", "Prepaid Insurance", "='Unadjusted TB'!B8/'Unadjusted TB'!B25*'Unadjusted TB'!B26"),
        ("Depreciation Expense", "Accumulated Depreciation", "='Unadjusted TB'!B27"),
        ("Wages Expense", "Wages Payable", "='Unadjusted TB'!B28"),
        ("Unearned Revenue", "Service Revenue", "='Unadjusted TB'!B29"),
        ("Accounts Receivable", "Service Revenue", "='Unadjusted TB'!B30"),
    )
    for row, (debit, credit, amount) in zip(range(5, 17, 2), adjustments):
        journal.cell(row, 3, debit)
        journal.cell(row, 4, amount)
        journal.cell(row, 5, 0)
        journal.cell(row + 1, 3, credit)
        journal.cell(row + 1, 4, 0)
        journal.cell(row + 1, 5, f"=D{row}")

    trial_balance = workbook["Adjusted TB"]
    for row in range(5, 20):
        trial_balance[f"D{row}"] = f'=SUMIF(Adjustments!$C$5:$C$16,A{row},Adjustments!$D$5:$D$16)'
        trial_balance[f"E{row}"] = f'=SUMIF(Adjustments!$C$5:$C$16,A{row},Adjustments!$E$5:$E$16)'
        trial_balance[f"F{row}"] = f"=MAX(B{row}+D{row}-C{row}-E{row},0)"
        trial_balance[f"G{row}"] = f"=MAX(C{row}+E{row}-B{row}-D{row},0)"

    statements = workbook["Financial Statements"]
    formulas = {
        "B5": "='Adjusted TB'!G15", "B6": "='Adjusted TB'!F16",
        "B7": "='Adjusted TB'!F17", "B8": "='Adjusted TB'!F18",
        "B9": "='Adjusted TB'!F19", "B10": "=SUM(B6:B9)", "B11": "=B5-B10",
        "B15": "=B11", "B16": "=B14+B15",
        "B19": "='Adjusted TB'!F5", "B20": "='Adjusted TB'!F6",
        "B21": "='Adjusted TB'!F7", "B22": "='Adjusted TB'!F8",
        "B23": "='Adjusted TB'!F9", "B24": "='Adjusted TB'!G10",
        "B25": "=SUM(B19:B23)-B24", "B27": "='Adjusted TB'!G11",
        "B28": "='Adjusted TB'!G12", "B29": "='Adjusted TB'!G13",
        "B30": "=SUM(B27:B29)", "B32": "=B16", "B33": "=B30+B32",
    }
    for address, formula in formulas.items():
        statements[address] = formula

    closing = workbook["Closing Entries"]
    entries = (
        ("Service Revenue", "='Financial Statements'!B5", 0),
        ("Income Summary", 0, "=D5"),
        ("Income Summary", "='Financial Statements'!B10", 0),
        ("Supplies Expense", 0, "='Financial Statements'!B6"),
        ("Insurance Expense", 0, "='Financial Statements'!B7"),
        ("Depreciation Expense", 0, "='Financial Statements'!B8"),
        ("Wages Expense", 0, "='Financial Statements'!B9"),
        ("Income Summary", "='Financial Statements'!B11", 0),
        ("Retained Earnings", 0, "=D12"),
    )
    for row, (account, debit, credit) in enumerate(entries, start=5):
        closing.cell(row, 3, account)
        closing.cell(row, 4, debit)
        closing.cell(row, 5, credit)

    post_closing = workbook["Post-closing TB"]
    for row in range(5, 15):
        post_closing[f"B{row}"] = f"='Adjusted TB'!F{row}"
        post_closing[f"C{row}"] = (
            "='Financial Statements'!B16" if row == 14 else f"='Adjusted TB'!G{row}"
        )


def main() -> None:
    """Read the confirmed source and update only the two website copies."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="Existing BM U02L04 - Textbook Aligned.xlsx")
    source = parser.parse_args().source.resolve()
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    student = load_workbook(source)
    if student.sheetnames != SHEETS or student["Unadjusted TB"]["B30"].value != 900:
        raise ValueError("The source does not match the confirmed six-sheet Lesson 4 worksheet.")
    prepare_website_copy(student)
    student.save(RESOURCES / "unit02-lesson04-student.xlsx")
    teacher = load_workbook(source)
    prepare_website_copy(teacher)
    complete_reference(teacher)
    teacher.save(RESOURCES / "unit02-lesson04-teacher.xlsx")
    if hashlib.sha256(source.read_bytes()).hexdigest() != source_hash:
        raise AssertionError("The OneDrive source changed during the copy.")
    print(f"Updated the two website copies. OneDrive source is unchanged: {source_hash}")


if __name__ == "__main__":
    main()
