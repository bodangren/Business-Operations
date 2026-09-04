"""Verify the Unit 1 workbook and dataset contracts."""

from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RESOURCES = ROOT / "public" / "resources"
PROJECT_SHEETS = ["Transactions", "Trial Balance", "Error Checks", "Executive Summary"]
HEADERS = ["Transaction ID", "Date", "Description", "Account", "Type", "Debit", "Credit"]


def require(condition: bool, message: str) -> None:
    """Raise an assertion error when a workbook contract fails."""
    if not condition:
        raise AssertionError(message)


def verify_shared_workbook(filename: str, expected_sheets: list[str]) -> None:
    """Verify sheet order, table structure, formulas, and shared transaction totals."""
    workbook = load_workbook(RESOURCES / filename, data_only=False)
    require(workbook.sheetnames == expected_sheets, f"{filename}: wrong sheet set")
    transactions = workbook["Transactions"]
    require([transactions.cell(1, column).value for column in range(1, 8)] == HEADERS, f"{filename}: wrong headers")

    if filename != "unit01-lesson04-student.xlsx":
        require("LedgerTable" in transactions.tables, f"{filename}: LedgerTable is missing")

    if transactions.max_row == 23:
        debits = sum(float(transactions.cell(row, 6).value or 0) for row in range(2, 24))
        credits = sum(float(transactions.cell(row, 7).value or 0) for row in range(2, 24))
        require(abs(debits - 8_620.98) < 0.001, f"{filename}: wrong debit total")
        require(abs(credits - 8_620.98) < 0.001, f"{filename}: wrong credit total")
        ids: dict[str, list[tuple[float, float]]] = defaultdict(list)
        for row in range(2, 24):
            ids[str(transactions.cell(row, 1).value)].append(
                (float(transactions.cell(row, 6).value or 0), float(transactions.cell(row, 7).value or 0))
            )
        require(len(ids) == 11, f"{filename}: expected 11 transaction IDs")
        require(all(len(rows) == 2 for rows in ids.values()), f"{filename}: each ID needs two rows")
        require(all(abs(sum(debit for debit, _ in rows) - sum(credit for _, credit in rows)) < 0.001 for rows in ids.values()), f"{filename}: unbalanced transaction")

    if "Trial Balance" in workbook.sheetnames:
        trial_balance = workbook["Trial Balance"]
        if filename == "unit01-lesson05-student.xlsx":
            require(trial_balance["B2"].value is None, f"{filename}: Trial Balance must be a blank template")
        else:
            require(trial_balance["B2"].value == "=SUMIF(LedgerTable[Account],A2,LedgerTable[Debit])", f"{filename}: wrong SUMIF formula")
            require(trial_balance["D2"].value == "=MAX(B2-C2,0)", f"{filename}: wrong debit-balance formula")

    if "Error Checks" in workbook.sheetnames:
        checks = workbook["Error Checks"]
        require(checks.max_row == 5, f"{filename}: expected four controls")
        if filename == "unit01-lesson05-student.xlsx":
            require(checks["B3"].value is None, f"{filename}: Error Checks must be a blank template")
        else:
            require(checks["B3"].value == '=COUNTIF(LedgerTable[Entry Difference],"<>0")', f"{filename}: wrong transaction control")

    if "Executive Summary" in workbook.sheetnames:
        summary = workbook["Executive Summary"]
        if filename in {"unit01-lesson05-checkpoint.xlsx", "unit01-lesson06-student.xlsx"}:
            require(summary["B7"].value is None, f"{filename}: Executive Summary must be a blank template")
        else:
            require(summary["B7"].value == '=COUNTIF(\'Error Checks\'!C2:C5,"Review")', f"{filename}: wrong review count")
            require(summary["B10"].value == "=B8-B9", f"{filename}: wrong net-income formula")


def verify_group_files(group_number: int, csv_name: str) -> None:
    """Verify one group workbook and its 12-event source dataset."""
    workbook_name = f"unit01-group{group_number}-starter.xlsx"
    verify_shared_workbook(workbook_name, PROJECT_SHEETS)
    with (RESOURCES / csv_name).open(newline="", encoding="utf-8") as source:
        rows = list(csv.DictReader(source))
    require(len(rows) == 12, f"{csv_name}: expected 12 events")
    require(list(rows[0]) == ["date", "transaction_id", "description", "debit_account", "credit_account", "amount"], f"{csv_name}: wrong headers")
    require(len({row["transaction_id"] for row in rows}) == 12, f"{csv_name}: transaction IDs must be unique")
    require(all(float(row["amount"]) > 0 for row in rows), f"{csv_name}: amounts must be positive")


def main() -> None:
    """Run all Unit 1 workbook and dataset checks."""
    verify_shared_workbook("unit01-lesson04-student.xlsx", ["Transactions"])
    verify_shared_workbook("unit01-lesson04-teacher.xlsx", ["Transactions"])
    verify_shared_workbook("unit01-lesson05-student.xlsx", ["Transactions", "Trial Balance", "Error Checks"])
    verify_shared_workbook("unit01-lesson05-teacher.xlsx", ["Transactions", "Trial Balance", "Error Checks"])
    verify_shared_workbook("unit01-lesson05-checkpoint.xlsx", PROJECT_SHEETS)
    verify_shared_workbook("unit01-lesson06-student.xlsx", PROJECT_SHEETS)
    verify_shared_workbook("unit01-lesson06-teacher.xlsx", PROJECT_SHEETS)
    verify_shared_workbook("unit01-lesson07-student.xlsx", PROJECT_SHEETS)
    verify_shared_workbook("unit01-lesson07-teacher.xlsx", PROJECT_SHEETS)
    verify_shared_workbook("unit01-rehearsal-workbook.xlsx", PROJECT_SHEETS)
    verify_group_files(1, "unit01-group1-foodtruck.csv")
    verify_group_files(2, "unit01-group2-ecommerce.csv")
    verify_group_files(3, "unit01-group3-tutoring.csv")
    verify_group_files(4, "unit01-group4-custom.csv")
    print("Unit 1 workbook verification passed.")


if __name__ == "__main__":
    main()
