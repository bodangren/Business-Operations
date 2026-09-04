"""Build the Unit 1 lesson and project workbook resources."""

from __future__ import annotations

import csv
from copy import copy
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
RESOURCES = ROOT / "public" / "resources"

PROJECT_SHEETS = ("Transactions", "Trial Balance", "Error Checks", "Executive Summary")
ACCOUNT_TYPES = {
    "Cash": "Asset",
    "Accounts Receivable": "Asset",
    "Supplies Inventory": "Asset",
    "Inventory": "Asset",
    "Equipment": "Asset",
    "Accounts Payable": "Liability",
    "Notes Payable": "Liability",
    "Unearned Revenue": "Liability",
    "Owner's Capital": "Equity",
    "Owner's Draw": "Contra-Equity",
    "Service Revenue": "Revenue",
    "Sales Revenue": "Revenue",
    "Advertising Expense": "Expense",
    "Software Subscription Expense": "Expense",
    "Software Expense": "Expense",
    "Rent Expense": "Expense",
    "Salaries Expense": "Expense",
    "Wages Expense": "Expense",
    "Supplies Expense": "Expense",
    "Cost of Goods Sold": "Expense",
    "Fuel Expense": "Expense",
    "Utilities Expense": "Expense",
}

TRIAL_BALANCE_ACCOUNTS = tuple(ACCOUNT_TYPES)
MONEY_FORMAT = '$#,##0.00;[Red]($#,##0.00);-'
THIN_GRAY = Side(style="thin", color="D1D5DB")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
INPUT_FILL = PatternFill("solid", fgColor="DDEBF7")
FORMULA_FILL = PatternFill("solid", fgColor="F3F4F6")
WARNING_FILL = PatternFill("solid", fgColor="FCE8E6")
PASS_FILL = PatternFill("solid", fgColor="E2F0D9")


@dataclass(frozen=True)
class Transaction:
    transaction_id: str
    transaction_date: date
    description: str
    debit_account: str
    credit_account: str
    amount: float


SHARED_TRANSACTIONS = (
    Transaction("T001", date(2025, 1, 3), "Bakery website payment received", "Cash", "Service Revenue", 2_200),
    Transaction("T002", date(2025, 1, 6), "Pet grooming social-media project billed", "Accounts Receivable", "Service Revenue", 650),
    Transaction("T003", date(2025, 1, 9), "Advertising campaign paid", "Advertising Expense", "Cash", 400),
    Transaction("T004", date(2025, 1, 12), "Adobe Creative Cloud subscription bought on account", "Software Subscription Expense", "Accounts Payable", 52.99),
    Transaction("T005", date(2025, 1, 15), "Dental office SEO payment received", "Cash", "Service Revenue", 1_100),
    Transaction("T006", date(2025, 1, 18), "Adobe subscription balance paid", "Accounts Payable", "Cash", 52.99),
    Transaction("T007", date(2025, 1, 21), "Office rent paid", "Rent Expense", "Cash", 850),
    Transaction("T008", date(2025, 1, 24), "Contractor wages paid", "Salaries Expense", "Cash", 1_200),
    Transaction("T009", date(2025, 1, 27), "Supplies used for client work", "Supplies Expense", "Supplies Inventory", 315),
    Transaction("T010", date(2025, 1, 29), "Additional client work billed", "Accounts Receivable", "Service Revenue", 900),
    Transaction("T011", date(2025, 1, 30), "Client invoice collected", "Cash", "Accounts Receivable", 900),
)

GROUP_TRANSACTIONS = {
    "group1": (
        "Food Truck Venture",
        "unit01-group1-foodtruck.csv",
        (
            Transaction("F001", date(2025, 2, 1), "Owner startup investment", "Cash", "Owner's Capital", 6_000),
            Transaction("F002", date(2025, 2, 2), "Cooking equipment purchased for cash", "Equipment", "Cash", 2_500),
            Transaction("F003", date(2025, 2, 3), "Food inventory bought on account", "Inventory", "Accounts Payable", 800),
            Transaction("F004", date(2025, 2, 5), "Lunch service cash sales", "Cash", "Sales Revenue", 1_500),
            Transaction("F005", date(2025, 2, 5), "Inventory used for lunch service", "Cost of Goods Sold", "Inventory", 520),
            Transaction("F006", date(2025, 2, 6), "Fuel paid in cash", "Fuel Expense", "Cash", 180),
            Transaction("F007", date(2025, 2, 8), "Catering invoice issued", "Accounts Receivable", "Sales Revenue", 1_200),
            Transaction("F008", date(2025, 2, 10), "Supplier balance paid", "Accounts Payable", "Cash", 800),
            Transaction("F009", date(2025, 2, 12), "Catering invoice collected", "Cash", "Accounts Receivable", 1_200),
            Transaction("F010", date(2025, 2, 14), "Crew wages paid", "Wages Expense", "Cash", 600),
            Transaction("F011", date(2025, 2, 16), "Weekend cash sales", "Cash", "Sales Revenue", 1_800),
            Transaction("F012", date(2025, 2, 16), "Inventory used for weekend service", "Cost of Goods Sold", "Inventory", 630),
        ),
    ),
    "group2": (
        "E-commerce Business",
        "unit01-group2-ecommerce.csv",
        (
            Transaction("E001", date(2025, 3, 1), "Owner startup investment", "Cash", "Owner's Capital", 7_500),
            Transaction("E002", date(2025, 3, 2), "Opening inventory bought for cash", "Inventory", "Cash", 2_200),
            Transaction("E003", date(2025, 3, 3), "Store software subscription paid", "Software Expense", "Cash", 120),
            Transaction("E004", date(2025, 3, 5), "Online sales collected", "Cash", "Sales Revenue", 1_900),
            Transaction("E005", date(2025, 3, 5), "Cost of products sold", "Cost of Goods Sold", "Inventory", 760),
            Transaction("E006", date(2025, 3, 7), "Digital advertising paid", "Advertising Expense", "Cash", 350),
            Transaction("E007", date(2025, 3, 9), "Wholesale order billed", "Accounts Receivable", "Sales Revenue", 1_400),
            Transaction("E008", date(2025, 3, 10), "New inventory bought on account", "Inventory", "Accounts Payable", 1_100),
            Transaction("E009", date(2025, 3, 12), "Wholesale invoice collected", "Cash", "Accounts Receivable", 1_400),
            Transaction("E010", date(2025, 3, 14), "Supplier balance paid", "Accounts Payable", "Cash", 1_100),
            Transaction("E011", date(2025, 3, 16), "Weekend online sales collected", "Cash", "Sales Revenue", 2_100),
            Transaction("E012", date(2025, 3, 16), "Cost of weekend products sold", "Cost of Goods Sold", "Inventory", 840),
        ),
    ),
    "group3": (
        "Tutoring Service",
        "unit01-group3-tutoring.csv",
        (
            Transaction("T001", date(2025, 4, 1), "Owner startup investment", "Cash", "Owner's Capital", 4_000),
            Transaction("T002", date(2025, 4, 2), "Teaching equipment purchased", "Equipment", "Cash", 900),
            Transaction("T003", date(2025, 4, 3), "Monthly room rent paid", "Rent Expense", "Cash", 650),
            Transaction("T004", date(2025, 4, 5), "Private lessons paid in cash", "Cash", "Service Revenue", 1_200),
            Transaction("T005", date(2025, 4, 6), "School workshop billed", "Accounts Receivable", "Service Revenue", 1_500),
            Transaction("T006", date(2025, 4, 8), "Tutor wages paid", "Wages Expense", "Cash", 700),
            Transaction("T007", date(2025, 4, 10), "Scheduling software paid", "Software Expense", "Cash", 80),
            Transaction("T008", date(2025, 4, 12), "School workshop invoice collected", "Cash", "Accounts Receivable", 1_500),
            Transaction("T009", date(2025, 4, 14), "Future lesson package deposit received", "Cash", "Unearned Revenue", 600),
            Transaction("T010", date(2025, 4, 16), "Private lessons paid in cash", "Cash", "Service Revenue", 1_350),
            Transaction("T011", date(2025, 4, 18), "Advertising paid", "Advertising Expense", "Cash", 200),
            Transaction("T012", date(2025, 4, 20), "Tutor wages paid", "Wages Expense", "Cash", 750),
        ),
    ),
    "group4": (
        "Event Planning Studio",
        "unit01-group4-custom.csv",
        (
            Transaction("P001", date(2025, 5, 1), "Owner startup investment", "Cash", "Owner's Capital", 5_500),
            Transaction("P002", date(2025, 5, 2), "Planning equipment purchased", "Equipment", "Cash", 1_200),
            Transaction("P003", date(2025, 5, 3), "Client deposit received for future event", "Cash", "Unearned Revenue", 1_000),
            Transaction("P004", date(2025, 5, 5), "Venue planning fee billed", "Accounts Receivable", "Service Revenue", 1_800),
            Transaction("P005", date(2025, 5, 6), "Design software paid", "Software Expense", "Cash", 95),
            Transaction("P006", date(2025, 5, 8), "Assistant wages paid", "Wages Expense", "Cash", 650),
            Transaction("P007", date(2025, 5, 10), "Venue planning invoice collected", "Cash", "Accounts Receivable", 1_800),
            Transaction("P008", date(2025, 5, 12), "Event service completed from deposit", "Unearned Revenue", "Service Revenue", 1_000),
            Transaction("P009", date(2025, 5, 14), "Studio rent paid", "Rent Expense", "Cash", 700),
            Transaction("P010", date(2025, 5, 16), "New event package paid in cash", "Cash", "Service Revenue", 1_500),
            Transaction("P011", date(2025, 5, 18), "Advertising paid", "Advertising Expense", "Cash", 250),
            Transaction("P012", date(2025, 5, 20), "Utilities paid", "Utilities Expense", "Cash", 140),
        ),
    ),
}


def style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.fill = copy(HEADER_FILL)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = Border(bottom=THIN_GRAY)


def set_widths(ws, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def set_print_layout(ws) -> None:
    """Fit the used columns to one landscape page width for review or printing."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = ws.calculate_dimension()
    ws.oddFooter.center.text = "Page &P of &N"


def journal_rows(transactions: tuple[Transaction, ...]) -> list[list[object]]:
    rows: list[list[object]] = []
    for transaction in transactions:
        rows.append([
            transaction.transaction_id,
            transaction.transaction_date,
            transaction.description,
            transaction.debit_account,
            ACCOUNT_TYPES[transaction.debit_account],
            transaction.amount,
            0,
        ])
        rows.append([
            transaction.transaction_id,
            transaction.transaction_date,
            transaction.description,
            transaction.credit_account,
            ACCOUNT_TYPES[transaction.credit_account],
            0,
            transaction.amount,
        ])
    return rows


def add_transactions_sheet(
    wb: Workbook,
    transactions: tuple[Transaction, ...] | None,
    *,
    table: bool,
    entry_check: bool,
) -> None:
    ws = wb.create_sheet("Transactions")
    headers = ["Transaction ID", "Date", "Description", "Account", "Type", "Debit", "Credit"]
    if entry_check:
        headers.append("Entry Difference")
    ws.append(headers)

    rows = journal_rows(transactions) if transactions else [["", "", "", "", "", 0, 0]]
    for row in rows:
        ws.append(row)

    if entry_check:
        for row_number in range(2, ws.max_row + 1):
            ws.cell(row_number, 8).value = (
                '=SUMIFS(LedgerTable[Debit],LedgerTable[Transaction ID],[@[Transaction ID]])'
                '-SUMIFS(LedgerTable[Credit],LedgerTable[Transaction ID],[@[Transaction ID]])'
            )

    style_header(ws)
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    set_widths(ws, {"A": 16, "B": 13, "C": 48, "D": 30, "E": 18, "F": 14, "G": 14, "H": 18})
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "yyyy-mm-dd"
        row[5].number_format = MONEY_FORMAT
        row[6].number_format = MONEY_FORMAT
        if entry_check:
            row[7].number_format = MONEY_FORMAT

    if table:
        end_column = "H" if entry_check else "G"
        ledger_table = Table(displayName="LedgerTable", ref=f"A1:{end_column}{ws.max_row}")
        ledger_table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(ledger_table)
    set_print_layout(ws)


def add_trial_balance_sheet(wb: Workbook, *, formulas: bool) -> int:
    ws = wb.create_sheet("Trial Balance")
    ws.append(["Account", "Total Debits", "Total Credits", "Debit Balance", "Credit Balance"])
    for account in TRIAL_BALANCE_ACCOUNTS:
        ws.append([account])
    total_row = ws.max_row + 1
    ws.cell(total_row, 1).value = "Totals"
    status_row = total_row + 2
    ws.cell(status_row, 1).value = "Difference"
    ws.cell(status_row + 1, 1).value = "Status"

    if formulas:
        for row_number in range(2, total_row):
            ws.cell(row_number, 2).value = f"=SUMIF(LedgerTable[Account],A{row_number},LedgerTable[Debit])"
            ws.cell(row_number, 3).value = f"=SUMIF(LedgerTable[Account],A{row_number},LedgerTable[Credit])"
            ws.cell(row_number, 4).value = f"=MAX(B{row_number}-C{row_number},0)"
            ws.cell(row_number, 5).value = f"=MAX(C{row_number}-B{row_number},0)"
        for column in range(2, 6):
            letter = ws.cell(1, column).column_letter
            ws.cell(total_row, column).value = f"=SUM({letter}2:{letter}{total_row - 1})"
        ws.cell(status_row, 2).value = f"=D{total_row}-E{total_row}"
        ws.cell(status_row + 1, 2).value = f'=IF(ABS(B{status_row})<0.01,"Balanced","Review Needed")'

    style_header(ws)
    for cell in ws[total_row]:
        cell.font = Font(bold=True)
        cell.fill = copy(SUBHEADER_FILL)
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=5):
        for cell in row:
            cell.number_format = MONEY_FORMAT
            if cell.value is not None:
                cell.fill = copy(FORMULA_FILL)
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    set_widths(ws, {"A": 32, "B": 16, "C": 16, "D": 16, "E": 16})
    set_print_layout(ws)
    return total_row


def add_error_checks_sheet(wb: Workbook, *, formulas: bool) -> None:
    ws = wb.create_sheet("Error Checks")
    ws.append(["Control", "Result", "Status", "Action if Review"])
    checks = [
        ("Ledger debit-credit difference", '=SUM(LedgerTable[Debit])-SUM(LedgerTable[Credit])', "Find a missing or unequal amount."),
        ("Rows in unbalanced transactions", '=COUNTIF(LedgerTable[Entry Difference],"<>0")', "Filter Entry Difference for nonzero values."),
        ("Rows with a blank account", '=COUNTBLANK(LedgerTable[Account])', "Complete each Account cell."),
        ("Rows with both or neither amount", '=SUMPRODUCT(--((LedgerTable[Debit]=0)=(LedgerTable[Credit]=0)))', "Each row needs one debit or one credit."),
    ]
    for row_number, (label, formula, action) in enumerate(checks, start=2):
        ws.cell(row_number, 1).value = label
        ws.cell(row_number, 4).value = action
        if formulas:
            ws.cell(row_number, 2).value = formula
            ws.cell(row_number, 3).value = f'=IF(ABS(B{row_number})<0.01,"Pass","Review")'

    style_header(ws)
    ws.sheet_view.showGridLines = False
    set_widths(ws, {"A": 36, "B": 18, "C": 14, "D": 48})
    for row in range(2, 6):
        ws.cell(row, 2).fill = copy(FORMULA_FILL)
        ws.cell(row, 3).fill = copy(PASS_FILL)
    ws.conditional_formatting.add("C2:C5", CellIsRule(operator="equal", formula=['"Review"'], fill=WARNING_FILL))
    set_print_layout(ws)


def add_summary_sheet(wb: Workbook, business_name: str, *, formulas: bool, shared: bool) -> None:
    ws = wb.create_sheet("Executive Summary")
    ws.append([f"{business_name} Smart Ledger Summary", ""])
    ws.append(["Prepared By", "Student Team"])
    labels = [
        "Total Debit Balances",
        "Total Credit Balances",
        "Difference",
        "Ledger Status",
        "Controls Needing Review",
        "Revenue",
        "Expenses",
        "Net Income",
        "Cash Balance",
        "Recommendation",
        "Risk or Limitation",
    ]
    for label in labels:
        ws.append([label, ""])

    if formulas:
        total_row = len(TRIAL_BALANCE_ACCOUNTS) + 2
        difference_row = total_row + 2
        ws["B3"] = f"='Trial Balance'!D{total_row}"
        ws["B4"] = f"='Trial Balance'!E{total_row}"
        ws["B5"] = f"='Trial Balance'!B{difference_row}"
        ws["B6"] = f"='Trial Balance'!B{difference_row + 1}"
        ws["B7"] = '=COUNTIF(\'Error Checks\'!C2:C5,"Review")'
        ws["B8"] = '=SUMIFS(LedgerTable[Credit],LedgerTable[Type],"Revenue")-SUMIFS(LedgerTable[Debit],LedgerTable[Type],"Revenue")'
        ws["B9"] = '=SUMIFS(LedgerTable[Debit],LedgerTable[Type],"Expense")-SUMIFS(LedgerTable[Credit],LedgerTable[Type],"Expense")'
        ws["B10"] = "=B8-B9"
        ws["B11"] = '=SUMIF(LedgerTable[Account],"Cash",LedgerTable[Debit])-SUMIF(LedgerTable[Account],"Cash",LedgerTable[Credit])'
        if shared:
            ws["B12"] = "TechStart can present the ledger because the controls pass and the business earned positive net income."
            ws["B13"] = "The ledger covers one month. It does not prove that future revenue or cash flow will remain stable."
        else:
            ws["B12"] = "Write one claim supported by at least three workbook values."
            ws["B13"] = "State one risk or limitation."

    ws.merge_cells("A1:B1")
    ws["A1"].fill = copy(HEADER_FILL)
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    for row_number in range(2, 14):
        ws.cell(row_number, 1).font = Font(bold=True)
        ws.cell(row_number, 1).fill = copy(SUBHEADER_FILL)
        ws.cell(row_number, 2).fill = copy(FORMULA_FILL if row_number < 12 else INPUT_FILL)
    for row_number in (3, 4, 5, 8, 9, 10, 11):
        ws.cell(row_number, 2).number_format = MONEY_FORMAT
    ws["B12"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["B13"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[12].height = 48
    ws.row_dimensions[13].height = 48
    ws.sheet_view.showGridLines = False
    set_widths(ws, {"A": 30, "B": 86})
    set_print_layout(ws)


def new_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    return wb


def build_lesson_workbooks() -> None:
    workbook = new_workbook()
    add_transactions_sheet(workbook, SHARED_TRANSACTIONS, table=False, entry_check=False)
    workbook.save(RESOURCES / "unit01-lesson04-student.xlsx")

    workbook = new_workbook()
    add_transactions_sheet(workbook, SHARED_TRANSACTIONS, table=True, entry_check=False)
    workbook.save(RESOURCES / "unit01-lesson04-teacher.xlsx")

    workbook = new_workbook()
    add_transactions_sheet(workbook, SHARED_TRANSACTIONS, table=True, entry_check=False)
    add_trial_balance_sheet(workbook, formulas=False)
    add_error_checks_sheet(workbook, formulas=False)
    workbook.save(RESOURCES / "unit01-lesson05-student.xlsx")

    workbook = new_workbook()
    add_transactions_sheet(workbook, SHARED_TRANSACTIONS, table=True, entry_check=True)
    add_trial_balance_sheet(workbook, formulas=True)
    add_error_checks_sheet(workbook, formulas=True)
    workbook.save(RESOURCES / "unit01-lesson05-teacher.xlsx")

    workbook = new_workbook()
    add_transactions_sheet(workbook, SHARED_TRANSACTIONS, table=True, entry_check=True)
    add_trial_balance_sheet(workbook, formulas=True)
    add_error_checks_sheet(workbook, formulas=True)
    add_summary_sheet(workbook, "TechStart Solutions", formulas=False, shared=True)
    workbook.save(RESOURCES / "unit01-lesson06-student.xlsx")
    workbook.save(RESOURCES / "unit01-lesson05-checkpoint.xlsx")

    workbook = new_workbook()
    add_transactions_sheet(workbook, SHARED_TRANSACTIONS, table=True, entry_check=True)
    add_trial_balance_sheet(workbook, formulas=True)
    add_error_checks_sheet(workbook, formulas=True)
    add_summary_sheet(workbook, "TechStart Solutions", formulas=True, shared=True)
    workbook.save(RESOURCES / "unit01-lesson06-teacher.xlsx")
    workbook.save(RESOURCES / "unit01-lesson07-student.xlsx")
    workbook.save(RESOURCES / "unit01-lesson07-teacher.xlsx")
    workbook.save(RESOURCES / "unit01-rehearsal-workbook.xlsx")


def write_group_dataset(filename: str, transactions: tuple[Transaction, ...]) -> None:
    path = RESOURCES / filename
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(["date", "transaction_id", "description", "debit_account", "credit_account", "amount"])
        for transaction in transactions:
            writer.writerow([
                transaction.transaction_date.isoformat(),
                transaction.transaction_id,
                transaction.description,
                transaction.debit_account,
                transaction.credit_account,
                f"{transaction.amount:.2f}",
            ])


def build_group_resources() -> None:
    for group_number, (_, (business_name, filename, transactions)) in enumerate(GROUP_TRANSACTIONS.items(), start=1):
        write_group_dataset(filename, transactions)
        workbook = new_workbook()
        add_transactions_sheet(workbook, None, table=True, entry_check=True)
        add_trial_balance_sheet(workbook, formulas=True)
        add_error_checks_sheet(workbook, formulas=True)
        add_summary_sheet(workbook, business_name, formulas=True, shared=False)
        workbook.save(RESOURCES / f"unit01-group{group_number}-starter.xlsx")


def validate_outputs() -> None:
    shared_total = round(sum(transaction.amount for transaction in SHARED_TRANSACTIONS), 2)
    assert shared_total == 8620.98
    assert all((RESOURCES / filename).exists() for filename in (
        "unit01-lesson04-student.xlsx",
        "unit01-lesson04-teacher.xlsx",
        "unit01-lesson05-student.xlsx",
        "unit01-lesson05-teacher.xlsx",
        "unit01-lesson06-student.xlsx",
        "unit01-lesson06-teacher.xlsx",
        "unit01-rehearsal-workbook.xlsx",
    ))


if __name__ == "__main__":
    build_lesson_workbooks()
    build_group_resources()
    validate_outputs()
