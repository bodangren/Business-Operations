"""Verify Excel for the web workbook structure and package safety."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RESOURCES = ROOT / "public" / "resources"
UNIT_TWO_FILES = (
    "unit02-lesson05-student.xlsx",
    "unit02-lesson05-teacher.xlsx",
    "unit02-lesson06-student.xlsx",
    "unit02-lesson06-teacher.xlsx",
)
FORBIDDEN_PARTS = ("vbaproject.bin", "xl/activex/", "xl/ctrlprops/")
FORBIDDEN_TEXT = re.compile(r"\b(?:vba|macros?)\b|\.xlsm\b|visual basic", re.IGNORECASE)


def require(condition: bool, message: str) -> None:
    """Raise an assertion error when a workbook contract fails."""
    if not condition:
        raise AssertionError(message)


def verify_package(path: Path) -> None:
    """Reject code projects, embedded controls, and forbidden workbook text."""
    require(path.suffix.lower() == ".xlsx", f"{path.name}: workbook must use .xlsx")
    with zipfile.ZipFile(path) as archive:
        names = [name.lower() for name in archive.namelist()]
        for forbidden in FORBIDDEN_PARTS:
            require(not any(forbidden in name for name in names), f"{path.name}: contains {forbidden}")

    workbook = load_workbook(path, data_only=False, read_only=True)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    require(
                        FORBIDDEN_TEXT.search(cell.value) is None,
                        f"{path.name}: forbidden text in {sheet.title}!{cell.coordinate}",
                    )


def verify_unit_two_contracts() -> None:
    """Verify workbook sheets, defined names, formulas, and controls."""
    for filename in UNIT_TWO_FILES:
        verify_package(RESOURCES / filename)

    lesson_five_student = load_workbook(RESOURCES / UNIT_TWO_FILES[0], data_only=False)
    require(lesson_five_student.sheetnames == ["Inputs", "Close Model", "Control Panel"], "Lesson 5 student: wrong sheets")
    require(lesson_five_student["Close Model"]["B4"].value is None, "Lesson 5 student: formula scaffold must be blank")

    lesson_five_teacher = load_workbook(RESOURCES / UNIT_TWO_FILES[1], data_only=False)
    require(lesson_five_teacher["Close Model"]["B4"].value == "=SuppliesUsed", "Lesson 5 teacher: wrong adjustment formula")
    require(lesson_five_teacher["Control Panel"]["B9"].value.startswith("=IF(AND("), "Lesson 5 teacher: CloseStatus is missing")

    lesson_six_student = load_workbook(RESOURCES / UNIT_TWO_FILES[2], data_only=False)
    require(lesson_six_student.sheetnames == ["Inputs", "Close Model", "Control Panel", "Scenarios"], "Lesson 6 student: wrong sheets")
    require(lesson_six_student["Inputs"]["B5"].value == 1_200, "Lesson 6 student: expected direct input scaffold")

    lesson_six_teacher = load_workbook(RESOURCES / UNIT_TWO_FILES[3], data_only=False)
    require(lesson_six_teacher["Inputs"]["B5"].value.startswith("=INDEX("), "Lesson 6 teacher: selected-period link is missing")
    require(lesson_six_teacher["Inputs"]["D5"].value.startswith("=IF(AND("), "Lesson 6 teacher: validation formula is missing")
    require(len(lesson_six_teacher["Control Panel"].data_validations.dataValidation) == 1, "Lesson 6 teacher: period dropdown is missing")
    require(lesson_six_teacher["Control Panel"]["B8"].value.startswith("=COUNTIF("), "Lesson 6 teacher: failed-check count is missing")

    expected_names = {
        "SuppliesUsed",
        "InsuranceExpired",
        "DepreciationExpense",
        "WagesAccrued",
        "RevenueEarned",
        "SelectedPeriod",
    }
    require(expected_names.issubset(set(lesson_six_teacher.defined_names)), "Lesson 6 teacher: defined names are missing")


def verify_all_resources() -> None:
    """Verify that every workbook resource is safe for Excel for the web."""
    for path in sorted(RESOURCES.glob("*.xlsx")):
        verify_package(path)


def main() -> None:
    """Run the Unit 2 contract checks and the repository workbook audit."""
    verify_unit_two_contracts()
    verify_all_resources()
    print("Excel for the web workbook verification passed.")


if __name__ == "__main__":
    main()
